"""
Genera docs/reunion_sandino_2026-09-03.docx con la transcripción de la reunión
de revisión del modelo de datos con el Dr. Sandino y el Dr. Israel Salas.

La transcripción vive aquí como dato para que el .docx sea reproducible y el
texto quede versionado. El análisis de lo que esta reunión cambia en el modelo
está en errores/fuentes/transcripcion_02_reunion_2026-09-03.md.

Ejecutar desde la raíz del repositorio:
    python scripts/gen_reunion_docx.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "docs" / "reunion_sandino_2026-09-03.docx"

ACENTO = RGBColor(0x0F, 0x6F, 0x70)
GRIS = RGBColor(0x67, 0x79, 0x78)

AF = "Ángel Frausto"
CS = "Dr. César Sandino"
IS = "Prof. Israel Salas"
VR = "Vanesa Rodríguez"

# (marca de tiempo, hablante, texto). None en el hablante = acotación.
TRANSCRIPCION = [
    ("00:00 – 00:32", AF,
     "…del modelo. Entonces, la primera y creo que la más importante —una disculpa si ya lo "
     "había preguntado, es que también perdimos el documento que teníamos de las anteriores "
     "preguntas—: en el protocolo que ya tienen en el laboratorio, ¿con qué se identifica "
     "físicamente a cada rata? O sea, ¿tienen algún arete numerado, marca, número de jaula, "
     "código de bitácora? ¿Y ese identificador es único en todo el laboratorio o solo dentro "
     "de un propio experimento?"),
    ("00:54 – 01:43", CS,
     "Regresan al bioterio. Lo que hacemos… la experimentación por tanda… Como son tratamientos "
     "que vamos… en el momento en el que vamos a hacer el… Las ratas las marcamos de la cola con "
     "un plumón indeleble y las marcamos con línea, ¿no?"),
    ("01:44 – 02:10", AF,
     "Ajá. ¿Y eso en el propio video tiene alguna identificación? O sea, tenemos el video con "
     "cuatro ratas en la imagen, ¿y se sabe que la rata en el cilindro número uno es tal rata, y "
     "en el cilindro número dos es tal rata, para que al día siguiente —en el caso de que sean "
     "dos días— también esté distribuido de la misma forma?"),
    ("02:11 – 02:17", CS,
     "Sí, a veces las ponemos en el mismo orden como se hacen en el primer video."),
    ("02:18 – 02:34", AF,
     "Okay. La siguiente pregunta sería: las ratas que salen juntas en un mismo video, ¿son "
     "siempre del mismo grupo experimental? ¿O alguna vez se graban juntas ratas de grupos "
     "distintos, o sea, de medicamento diferente?"),
    ("02:35 – 02:43", CS,
     "No, siempre utilizamos para cada video ratas que hayan recibido el mismo tratamiento."),
    ("02:44 – 03:13", AF,
     "Okay. Ahora la siguiente pregunta la tenemos escrita de esta forma: cuando la misma tanda "
     "se graba los dos días, ¿cada rata vuelve al mismo cilindro que ocupó el día uno o se "
     "reacomodan? Y no sé si en el protocolo que ustedes llevan en el laboratorio le llaman de "
     "esta forma; nosotros le pusimos esta palabra de «tanda». Pero tengo entendido que un "
     "experimento llega a tener más de cuatro ratas que no pueden ponerse en el mismo video, "
     "¿verdad?"),
    ("03:14 – 03:54", CS,
     "Así es. Normalmente los grupos son de 8 ratas para cada tratamiento. Le llamamos sesiones a "
     "los 5 minutos que se graban, o bueno, a la primera sesión de 20 minutos de nado, y luego a "
     "las 24 horas tenemos la sesión que ya es de captura de resultados. Siempre tenemos 8 ratas "
     "por grupo, entonces siempre vamos a tener como por lo menos dos videos por grupo con el "
     "mismo tratamiento."),
    ("03:55 – 04:10", AF,
     "Y en el… o sea, dos videos del mismo grupo, que serían para que se vean en la imagen cuatro "
     "ratas, ¿no? A eso es a lo que en nuestro lado le llamamos tanda. No sé si ahí también es a "
     "lo que se refiere con sesiones."),
    ("04:11 – 04:24", CS,
     "No, serían la mitad del grupo y la otra mitad. O sea, nada más las vamos identificando por "
     "el número de rata, que sería del 1 al 8."),
    ("04:25 – 04:55", AF,
     "Okay, entonces ahí nos atribuimos esa palabra para darle más sentido a lo que al ratito le "
     "vamos a explicar del modelo. Entonces, sabiendo esto, cuando la misma tanda —las cuatro "
     "primeras ratas que se grabaron de ocho— se grabarían igual en el mismo cilindro, ¿verdad? "
     "O sea, en el día uno y en el día dos."),
    ("04:53 – 04:55", CS, "Así es, sí."),
    ("04:56 – 05:07", AF,
     "Okay. La siguiente pregunta es: ¿el rango de 6 a 8 especímenes por grupo es fijo o alguna "
     "vez han trabajado con grupos más grandes?"),
    ("05:08 – 05:31", CS,
     "Depende del tipo de experimento que estemos haciendo. Comúnmente el máximo de ratas que "
     "llegamos a utilizar son 12, aunque con 6 ratas para muchos de los fármacos que probamos, "
     "con 6 o con 8 ratas es suficiente."),
    ("05:32 – 05:42", AF,
     "Okay. Entonces en el propio sistema, ¿podemos dejarle un límite en 12 o quiere que no tenga "
     "límite?"),
    ("05:43 – 06:03", CS,
     "Yo pensaría que puede ser que no tenga límite, porque incluso llegamos a repetir el "
     "experimento con otras ocho cuando no tenemos resultados que nos den una baja varianza entre "
     "la conducta de los primeros ocho."),
    ("06:04 – 06:09", AF, "Disculpe, creo que se me trabó un poquito. ¿Me podría responder, por favor?"),
    ("06:10 – 06:14", CS, "Yo preferiría que fuera sin límite."),
    ("06:13 – 06:14", AF, "¿Sin límite? Okay."),
    ("06:17 – 06:23", AF, "La siguiente pregunta: ¿los experimentos se cuentan por semestre o por bimestre?"),
    ("06:24 – 06:45", CS,
     "No, los experimentos son dependiendo del proyecto, cómo vaya avanzando. En un proyecto "
     "podemos tener 5 o 6 experimentos diferentes, y ese proyecto puede durar entre 1 año y 4 años."),
    ("06:46 – 07:16", AF,
     "Okay. Bueno, la siguiente pregunta sería: si un video se reanaliza y da un resultado "
     "ligeramente distinto, ¿se prefiere conservar ambos resultados para comparar o que el nuevo "
     "reemplace al anterior? Como tal, esto sería una varianza mínima, pero si se llega a analizar "
     "por segunda vez el mismo video, ¿le gustaría que conserváramos los resultados o no?"),
    ("07:17 – 07:29", CS,
     "Pero… ¿en qué circunstancia se tendría que reanalizar un video? ¿Cuál sería la causa para "
     "reanalizar un video?"),
    ("07:30 – 07:51", AF,
     "Bueno, como tal causa no habría más que solo… en algún punto diferente se quiere volver a "
     "subir el mismo video, o algún otro investigador lo sube de nuevo… un investigador diferente. "
     "Ahora sí que eso no…"),
    ("07:53 – 08:06", IS,
     "Lo que yo quiero entender que te está preguntando, es para ti, Frausto: ¿cuáles son los "
     "criterios para volver a hacer ese video?"),
    ("08:07 – 08:52", AF,
     "Bueno, en general, lo que el sistema estaría hablando es que si hay una mala iluminación, "
     "pues puede que incluso lo rechace, pero ya habría quedado un registro del medicamento y todo "
     "eso. Pero si el video con los mismos medicamentos se vuelve a modificar la iluminación… se "
     "va a analizar de nuevo con los mismos medicamentos o la misma tanda, pero un video mejorado, "
     "pero el análisis va a arrojar otra cosa. No sé si me doy a entender."),
    ("08:53 – 09:03", CS,
     "Sí. No, tendría que ser… Si hacemos alguna modificación, la que sea, nos quedaríamos con el "
     "análisis del video que estamos ejecutando al final."),
    ("09:05 – 09:38", AF,
     "Okay. Estoy checando mis apuntes… Creo que con esas preguntas por el momento nos respondió "
     "la mayoría. Ahora nada más serían tres cosas puntuales. La última vez que nos vimos nos había "
     "comentado que se estaban localizando los archivos de Excel donde los alumnos habían capturado "
     "los análisis manuales para empatar cada archivo que ya tenemos del video. No sé si los pudo "
     "ubicar."),
    ("10:00 – 10:48", CS,
     "Perdón, tenía apagado el micrófono. Creo que en Teams había un archivo de Excel, déjenme "
     "ahorita me asomo… Lo que pasa es que tenemos un detalle con esos videos: en algunos de esos "
     "videos se utilizó una técnica diferente para hacer la cuantificación. Hay dos técnicas: una "
     "que es la que quieren desarrollar y automatizar con su modelo, que es contabilizar por "
     "segundo lo que se está observando en las conductas. Hay otra técnica que se llama «por "
     "eventos»."),
    ("10:50 – 11:32", CS,
     "¿En qué consiste? Nosotros visualizamos el experimento y cada 5 segundos registramos la "
     "conducta que tiene. Si a los 5 segundos está nadando, nosotros lo ponemos como un evento de "
     "nado. Pero si a los otros 5 segundos —a los 10 segundos totales del video— está escalando, "
     "nosotros lo contabilizamos como un evento de escalamiento. Y si a los 15 segundos está "
     "nadando, exactamente a los 15 segundos, lo contabilizamos como un evento de nado. Y así "
     "hacemos la sumatoria de todos los eventos."),
    ("11:32 – 12:06", CS,
     "Esa técnica, que la utilizábamos cuando lo hacíamos totalmente a mano, era para tratar de "
     "estandarizar lo más posible los eventos que estábamos viendo. Tenemos unos archivos ya con "
     "los tiempos de la conducta… A ver, déjenme ver…"),
    (None, None, "(Pausa mientras el Dr. Sandino busca los archivos en su computadora)"),
    ("12:39 – 13:30", CS,
     "No están en las carpetas… Tengo un video que se llama «Resultados de videos», déjenme ver "
     "exactamente de qué es… Lo que no tendría ahorita es el empate entre qué videos de los que "
     "tienen corresponden a qué tiempos de los de la tabla de Excel. ¿Sí me explico? No sé qué "
     "tanta utilidad vaya a ser tenerlos así."),
    ("13:31 – 14:30", AF,
     "Con la explicación que nos acaba de dar podemos armar algo… La siguiente pregunta era que nos "
     "había comentado que tenía una guía que utilizaban los alumnos… un manual, si mal no recuerdo."),
    ("14:11 – 14:29", CS,
     "Era un video que les puse en el chat, ¿no? Creo que les puse en el chat un video de YouTube "
     "de las conductas. No sé si recuerdas, Israel, si por ahí andaba un video en el chat."),
    ("14:32 – 14:38", IS, "No, yo no detecté que hubiera un video de las conductas de las ratas."),
    ("14:39 – 14:45", CS, "¿Se acuerdan de la presentación?"),
    ("14:45 – 15:06", AF,
     "Sí, la que tenía tres videos de las ratitas, pero solamente funcionando, no había una "
     "explicación."),
    ("15:10 – 15:29", CS,
     "Te refieres a en qué momento o de qué manera se va a considerar escalamiento, nado o "
     "inmovilidad, ¿no?"),
    ("15:20 – 15:29", AF,
     "Sí, sobre todo para tener este manual para ponerlo en el documento como referencia completa."),
    ("15:30 – 15:39", CS, "Eso está descrito en el artículo de Porsolt. Ya tienen ese artículo, ¿no?"),
    ("15:39 – 16:08", AF,
     "Lo que pasa es que en Teams no tenemos… nos lo había pasado en un Drive. No nos había dado "
     "acceso al Teams."),
    ("16:10 – 16:35", CS, "¿No están con acceso al Teams?"),
    ("16:10 – 16:35", IS, "Yo que recuerde sí… A ver, comparte tu pantalla, Frausto, para que el doctor lo vea."),
    (None, None, "(Ángel comparte su pantalla mostrando Microsoft Teams y luego Google Drive)"),
    ("16:36 – 17:51", AF,
     "La carpeta que habíamos abierto y descargado había sido de Drive… Esta fue la carpeta que nos "
     "compartió: «videos_nado_forzado», fechadas en 2016."),
    ("17:59 – 19:27", CS, "¿En el Drive? ¿Con qué cuenta aparece ahí?"),
    ("17:59 – 19:27", AF, "Dice cesarsandino95."),
    ("17:59 – 19:27", CS, "Sí, esa es mi cuenta del Drive."),
    ("19:28 – 20:30", CS,
     "Pero no tengo… Ahorita nada más veo seis videos. Ahorita los uno. ¿Tu correo institucional?"),
    ("20:34 – 21:25", AF, "Es afraustor1500@alumno.ipn.mx."),
    ("20:34 – 21:25", CS, "¿Y Vanesa?"),
    ("20:34 – 21:25", VR, "Es vrodriguezverv2100@alumno.ipn.mx."),
    ("20:34 – 21:25", CS, "Listo. Ya están agregados en el grupo de Teams."),
    ("21:26 – 22:45", CS,
     "Ahí en general, en la pestaña compartidos, ahí están los videos… Déjenme ver…"),
    (None, None, "(El Dr. Sandino revisa sus carpetas en la nube)"),
    ("22:45 – 23:35", CS,
     "Aquí nada más tengo los de ansiedad, los de laberinto en cruz. ¿Cuántos videos tienen ahorita "
     "con los que han trabajado?"),
    ("22:45 – 23:35", AF, "Nosotros seis."),
    ("24:36 – 26:14", CS,
     "Esos los debo tener en la otra computadora, en la oficina. Eso lo subo mañana y voy a hacer "
     "una carpeta que va a decir «nado forzado»."),
    ("26:15 – 26:50", CS,
     "Hay un artículo donde justamente describen el modelo: cómo se interpreta, qué es lo que se "
     "evalúa en cada uno y qué se considera para cada conducta."),
    ("26:51 – 27:36", CS,
     "En YouTube hay un canal de protocolos, pero creo que ya no está libre. Lo voy a buscar a ver "
     "si tengo el video para que ahí vean cómo explican paso por paso qué conducta está haciendo y "
     "cómo debe interpretarse."),
    ("27:40 – 28:20", AF,
     "Quisiera ver si lo que llevamos ahorita de entendido del sistema va adecuado siguiendo este "
     "diagrama."),
    (None, None, "(Muestra en pantalla el diagrama conceptual entidad-relación)"),
    ("27:40 – 28:20", IS, "Si lo haces un poquito más grande… Sí, yo ya lo veo."),
    ("28:20 – 30:04", AF,
     "Lo que nosotros estamos haciendo con el sistema es: el usuario registra el experimento, "
     "empieza con el experimento y después sube los videos. El investigador va a separar los "
     "especímenes en grupos según lo que recibe cada uno: control (sin fármaco), referencia "
     "(antidepresivo) y tratamiento experimental. Manejando entre 6 y 8 ratas por grupo. La cámara "
     "web encuadra 4 cilindros a la vez, por lo que un grupo de 6 a 8 no cabe y se graba en dos "
     "tandas. 4 grupos por 8 ratas serían 32 especímenes."),
    ("30:04 – 31:00", AF,
     "El grupo se graba en tandas; la tanda va a alojar 3 a 4 especímenes y va a producir el video "
     "—uno o dos dependiendo de los días—. El espécimen aparece en el video y se procesa el análisis."),
    ("31:00 – 32:40", AF,
     "En cuanto a la identificación de las ratas, ¿se identifican con números naturales (1, 2, 3…)?"),
    ("31:00 – 32:40", CS,
     "Sí, por números. En el grupo es donde va identificado el tratamiento: 8, 6 o 12 ratas."),
    ("33:08 – 34:02", CS,
     "Nada más tengo una duda: ¿se van a analizar las 4 ratas al mismo tiempo en el video o se va a "
     "analizar rata por rata?"),
    ("33:08 – 34:02", AF,
     "Se va a analizar rata por rata; la rata 1 del video 1 se analiza y después se hace la "
     "comparación 1 a 1 con la rata en esa misma posición."),
    ("34:43 – 35:39", CS,
     "La información la vamos a obtener de la segunda sesión: comparar la conducta de todas las "
     "ratas de la segunda sesión. La primera sesión la tenemos como un control para ver la conducta "
     "de una rata que no esté estresada (20 minutos de nado). Esos valores de la primera sesión no "
     "nos sirven para la comparación del efecto del tratamiento; la que nos sirve es la segunda "
     "sesión (5 minutos)."),
    ("35:40 – 37:30", CS,
     "El experimento es así: tenemos tres grupos separados de 8 ratas cada uno. A todas esas ratas "
     "las ponemos 20 minutos a nadar el primer día (sin tratamiento). Al día siguiente les damos el "
     "tratamiento —placebo al control, fármaco estándar a la referencia, y tratamiento experimental "
     "al tercer grupo— y las grabamos 5 minutos. Esos 5 minutos del segundo día son los que se "
     "comparan entre los tres grupos."),
    ("37:31 – 38:29", CS,
     "El análisis de los 20 minutos del primer día es esencialmente para verificar que no haya "
     "ningún sesgo o artefacto en el bioterio (temperatura, ruido, etc.) y que todas partan de una "
     "conducta base similar."),
    ("40:08 – 41:28", IS,
     "César, sugiero que subas la presentación a Teams. Y Frausto, recuerda que todos los cambios "
     "que hagan en el modelo deben estar justificados y representados en el documento, porque los "
     "evaluadores o sinodales se los van a pedir."),
    ("41:29 – 42:40", CS,
     "Acabo de crear una carpeta en Teams que se llama «Sesión Depresión» dentro de la pestaña de "
     "Compartidos y acabo de subir la presentación. Si es posible, la próxima semana podríamos "
     "hacer la reunión presencial en ESCOM."),
    (None, None, "(Coordinación de horarios para la próxima reunión)"),
    ("42:44 – 45:30", CS,
     "En las tardes, entre 2:00 y 3:00 o 4:00 PM puedo ir a ESCOM, de preferencia lunes o viernes."),
    ("42:44 – 45:30", IS,
     "Para mí el lunes de 2:00 a 4:00 PM estaría perfecto, porque el martes lo tengo completamente "
     "lleno y el miércoles libre solo de 4:30 a 6:00 PM."),
    ("45:31 – 47:28", CS,
     "Quedamos en ver la disponibilidad de la Mtra. Martha para lunes o viernes de 2:00 a 4:00 PM. "
     "Por favor, mándenme un mensaje mañana en la mañana para recordarme subir los videos faltantes "
     "y los artículos. Muchas gracias a todos, nos vemos."),
]

RESUMEN = [
    ("Identificación del espécimen",
     "Las ratas se marcan en la cola con plumón indeleble, con líneas, y se numeran del 1 al 8 "
     "dentro de su grupo. El identificador es relativo al grupo, no único en el laboratorio."),
    ("Composición de los grupos",
     "8 ratas por grupo es lo normal; el rango real va de 6 a 12 y el Dr. Sandino pidió "
     "explícitamente que el sistema no imponga límite."),
    ("Una grabación no mezcla grupos",
     "Confirmado: cada video contiene únicamente ratas que recibieron el mismo tratamiento."),
    ("Posición entre sesiones",
     "Las ratas se colocan en el mismo orden en la segunda sesión que en la primera."),
    ("Cadencia de los experimentos",
     "No es por semestre ni por bimestre: depende del avance del proyecto. Un proyecto tiene 5 o 6 "
     "experimentos y dura de 1 a 4 años."),
    ("Reanálisis de un video",
     "Se conserva únicamente el análisis más reciente; el nuevo reemplaza al anterior."),
    ("Dos técnicas de cuantificación",
     "Además del conteo continuo por segundo que el sistema automatiza, el laboratorio ha usado una "
     "técnica «por eventos»: muestreo instantáneo cada 5 segundos. Los análisis manuales existentes "
     "están hechos con ambas técnicas."),
    ("Trazabilidad de los análisis manuales",
     "No existe el empate entre los archivos de Excel y los videos concretos a los que corresponden."),
    ("Propósito de la sesión del Día 1",
     "Los 20 minutos del primer día sirven para verificar que no haya sesgo ni artefacto en el "
     "bioterio y que todos los grupos partan de una conducta base similar. La comparación del "
     "efecto del tratamiento se hace con la segunda sesión."),
]

COMPROMISOS = [
    ("Dr. Sandino", "Subir a Teams los videos de nado forzado en una carpeta nueva llamada «nado forzado»", "Al día siguiente"),
    ("Dr. Sandino", "Compartir el artículo de Porsolt, donde se describe cómo se interpreta cada conducta", "Al día siguiente"),
    ("Dr. Sandino", "Buscar el video de protocolos que explica paso por paso cada conducta", "Sin fecha"),
    ("Dr. Sandino", "Presentación subida a Teams › Compartidos › «Sesión Depresión»", "Hecho en la sesión"),
    ("Dr. Sandino", "Acceso a Teams concedido a Ángel y Vanesa", "Hecho en la sesión"),
    ("Equipo", "Recordarle al Dr. Sandino por la mañana que suba videos y artículos", "Al día siguiente"),
    ("Equipo", "Justificar y representar en el documento todos los cambios del modelo (indicación del Prof. Salas)", "Continuo"),
    ("Todos", "Reunión presencial en ESCOM, lunes o viernes de 14:00 a 16:00, sujeta a la disponibilidad de la Mtra. Martha Cordero", "Próxima semana"),
]


def parrafo(doc, texto, size=10.5, bold=False, italic=False, color=None,
            space_after=6, space_before=0, align=None, izq=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if izq is not None:
        p.paragraph_format.left_indent = Cm(izq)
    if align is not None:
        p.alignment = align
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return p


def main():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.4)
        s.top_margin = s.bottom_margin = Cm(2.2)

    # ── portada ──
    parrafo(doc, "TT 2026-B066 · ESCOM-IPN", size=9, color=ACENTO, space_after=2)
    h = doc.add_heading("Reunión de revisión del modelo de datos", level=0)
    for r in h.runs:
        r.font.size = Pt(22)
    parrafo(doc, "Transcripción de la sesión con el laboratorio", size=12,
            italic=True, color=GRIS, space_after=14)

    ficha = doc.add_table(rows=0, cols=2)
    ficha.style = "Light List Accent 1"
    for k, v in [
        ("Fecha", "3 de septiembre de 2026"),
        ("Modalidad", "Reunión virtual"),
        ("Duración", "47 minutos"),
        ("Participantes",
         "Dr. César Augusto Sandino Reyes López (investigador colaborador, ENMyH-IPN) · "
         "Prof. Israel Salas Ramírez (director) · Ángel Ali Frausto Robles · "
         "Vanesa Rodríguez Verdín Sandoval"),
        ("Ausente", "Mtra. Martha Rosa Cordero López (directora) — la revisión de notación queda pendiente"),
        ("Propósito",
         "Verificar que el modelo de datos rediseñado corresponde al protocolo real del "
         "laboratorio, y cerrar las preguntas abiertas de diseño"),
    ]:
        fila = ficha.add_row().cells
        fila[0].width = Cm(3.4)
        fila[1].width = Cm(12.2)
        pk = fila[0].paragraphs[0]
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.size = Pt(9.5)
        pv = fila[1].paragraphs[0]
        rv = pv.add_run(v)
        rv.font.size = Pt(9.5)

    # ── resumen ──
    doc.add_heading("Lo que se resolvió", level=1)
    for titulo, texto in RESUMEN:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(titulo + ". ")
        r1.bold = True
        r1.font.size = Pt(10)
        r2 = p.add_run(texto)
        r2.font.size = Pt(10)

    # ── compromisos ──
    doc.add_heading("Acuerdos y compromisos", level=1)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    enc = t.rows[0].cells
    for i, (txt, ancho) in enumerate([("Responsable", 2.8), ("Compromiso", 9.6), ("Plazo", 3.2)]):
        enc[i].width = Cm(ancho)
        r = enc[i].paragraphs[0].add_run(txt)
        r.bold = True
        r.font.size = Pt(9)
    for quien, que, plazo in COMPROMISOS:
        fila = t.add_row().cells
        for i, (txt, ancho) in enumerate([(quien, 2.8), (que, 9.6), (plazo, 3.2)]):
            fila[i].width = Cm(ancho)
            r = fila[i].paragraphs[0].add_run(txt)
            r.font.size = Pt(9)

    # ── transcripción ──
    doc.add_heading("Transcripción cronológica", level=1)
    parrafo(doc,
            "Las marcas de tiempo provienen de la grabación. Algunos bloques se solapan: es un "
            "artefacto de la separación automática de hablantes, no una corrección posterior.",
            size=9, italic=True, color=GRIS, space_after=10)

    for marca, hablante, texto in TRANSCRIPCION:
        if hablante is None:
            parrafo(doc, texto, size=9.5, italic=True, color=GRIS,
                    space_before=4, space_after=6, izq=0.6)
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(0)
        rm = p.add_run("[%s]  " % marca)
        rm.font.size = Pt(8)
        rm.font.color.rgb = GRIS
        rh = p.add_run(hablante)
        rh.bold = True
        rh.font.size = Pt(9.5)
        rh.font.color.rgb = ACENTO
        parrafo(doc, texto, size=10.5, space_before=1, space_after=2, izq=0.6,
                align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # ── nota de cierre ──
    doc.add_page_break()
    doc.add_heading("Nota sobre esta transcripción", level=1)
    parrafo(doc,
            "El análisis de lo que esta reunión cambia en el modelo de datos, incluidas las "
            "contradicciones detectadas contra la entrevista formal previa, está en "
            "errores/fuentes/transcripcion_02_reunion_2026-09-03.md.",
            size=10, space_after=8)
    parrafo(doc,
            "Rango de fuente 3 (transcripción de reunión) según la jerarquía de fuentes de verdad "
            "del proyecto. Prevalece sobre el documento LaTeX, pero por debajo de las respuestas "
            "directas del equipo y de la entrevista formal del Anexo C.",
            size=10, italic=True, color=GRIS)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("Generado: %s" % OUT)


if __name__ == "__main__":
    main()
